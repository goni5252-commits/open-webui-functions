"""Offline Action tests with real XLSX writing and independent openpyxl reads."""
import asyncio
import contextlib
import importlib.util
import io
import json
import shutil
import subprocess
import sys
import threading
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch
from openpyxl import load_workbook

spec = importlib.util.spec_from_file_location('excel_export', Path(__file__).resolve().parents[1] / 'functions/엑셀로_내보내기.py')
e = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = e
spec.loader.exec_module(e)
TABLE = '| 항목 | 값 |\n|---|---|\n| 이름 | 가나다 |'

def workbook(action, markdown, names=None):
    tables = action.extract_tables_from_message(markdown)
    with contextlib.redirect_stdout(io.StringIO()):
        raw = action._build_workbook(tables, names or [f'Sheet{i+1}' for i in range(len(tables))])
    return load_workbook(io.BytesIO(raw), data_only=False)

class ExcelTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.a = e.Action()

    async def test_default_preserves_types_and_formulas(self):
        w = workbook(self.a, '| 값 |\n|---|\n|00123|\n|1234567890123456789|\n|2026-09-27|\n|=1+1|\n|https://example.org|\n|+123|')
        values = [w.active.cell(i,1).value for i in range(2,8)]
        self.assertEqual(values, ['00123','1234567890123456789','2026-09-27','=1+1','https://example.org','+123'])
        self.assertTrue(all(w.active.cell(i,1).data_type=='s' for i in range(2,8)))
        self.assertIsNone(w.active['A6'].hyperlink)

    async def test_safe_numbers_preserve_identifiers(self):
        self.a.valves.NUMBER_MODE='safe_numbers'
        w=workbook(self.a, '| 값 |\n|---|\n|42|\n|-3.5|\n|00123|\n|1234567890123456|\n|1e3|\n|-0.5|')
        self.assertEqual([w.active.cell(i,1).value for i in range(2,8)],[42,-3.5,'00123','1234567890123456','1e3','-0.5'])

    async def test_blank_duplicate_headers_and_extra_columns(self):
        w=workbook(self.a,'| A | | A |\n|---|---|---|\n| x | y | z | extra |\n| short |')
        self.assertEqual(list(w.active.values), [('A','Col2','A','Col4'),('x','y','z','extra'),('short',None,None,None)])

    async def test_escaped_and_code_pipes(self):
        w=workbook(self.a, '| A | B |\n|---|---|\n| x\\|y | `a|b` |')
        self.assertEqual(list(w.active.values)[1],('x|y','a|b'))

    async def test_fences_and_plain_pipes_not_tables(self):
        text='```markdown\n'+TABLE+'\n```\n\nplain | prose\n\n'+TABLE
        self.assertEqual(len(self.a.extract_tables_from_message(text)),1)
        self.assertEqual(self.a.extract_tables_from_message('a | b\nc | d'),[])
        self.assertEqual(self.a.extract_tables_from_message('    | a | b |\n    |---|---|'),[])

    async def test_optional_border_and_alignment(self):
        w=workbook(self.a,'A | B\n:--- | ---:\nleft | right')
        self.assertEqual(list(w.active.values),[('A','B'),('left','right')])

    async def test_case_long_reserved_sheet_names(self):
        names=['x'*31,'x'*31,'DATA','data',"'bad[]:*?/\\'",'History','']
        tables=[{'data':[['A'],['v']]} for _ in names]
        with contextlib.redirect_stdout(io.StringIO()):
            raw=self.a._build_workbook(tables,names)
        w=load_workbook(io.BytesIO(raw))
        self.assertEqual(len(w.sheetnames),7)
        self.assertTrue(all(len(n)<=31 for n in w.sheetnames))
        self.assertEqual(len({n.casefold() for n in w.sheetnames}),7)
        self.assertNotIn('History',w.sheetnames)

    async def test_oversized_cell_rejected_without_truncation(self):
        with self.assertRaisesRegex(ValueError,'32,767'):
            self.a._build_workbook([{'data':[['A'],['x'*32768]]}],['A'])

    async def test_total_cells_and_sheet_limit(self):
        self.a.valves.MAX_CELLS=2
        with self.assertRaises(ValueError):
            self.a._build_workbook([{'data':[['A','B'],['1','2']]}],['A'])
        self.a.valves.MAX_SHEETS=1
        with self.assertRaises(ValueError):
            self.a._build_workbook([{'data':[['A']]},{'data':[['A']]}],['A','B'])

    async def test_sheet_failure_aborts_whole_export(self):
        with patch.object(self.a,'apply_enhanced_formatting',side_effect=RuntimeError('write failed')), contextlib.redirect_stdout(io.StringIO()):
            with self.assertRaisesRegex(ValueError,'일부 표만'):
                self.a._build_workbook([{'data':[['A'],['v']]}],['A'])

    async def test_structured_output_excludes_reasoning(self):
        msg={'output':[{'type':'reasoning','content':TABLE},{'type':'message','content':[{'type':'output_text','text':TABLE}]}]}
        self.assertEqual(self.a._message_text(msg),TABLE)
        self.assertEqual(self.a._message_text({'content':[{'type':'text','text':TABLE}]}),TABLE)

    async def test_sync_async_db(self):
        self.assertEqual(await e._call_db(lambda:1),1)
        self.assertEqual(await e._call_db(AsyncMock(return_value=2)),2)

    async def test_owned_chat_recovers_output(self):
        chats=SimpleNamespace(get_chat_by_id_and_user_id=AsyncMock(return_value=SimpleNamespace(title='제목')))
        stored=SimpleNamespace(get_message_by_id=AsyncMock(return_value=SimpleNamespace(output=[{'type':'message','content':[{'type':'output_text','text':TABLE}]}])))
        with patch.object(e,'Chats',chats),patch.object(e,'ChatMessages',stored),contextlib.redirect_stdout(io.StringIO()):
            result=await self.a.action({'chat_id':'c','messages':[{'id':'m','content':''}]},{'id':'u'},AsyncMock(),AsyncMock(return_value={'ok':True}))
        self.assertIsNone(result)
        stored.get_message_by_id.assert_awaited_once_with('c-m')

    async def test_denied_chat_never_reads_message(self):
        chats=SimpleNamespace(get_chat_by_id_and_user_id=AsyncMock(return_value=None))
        stored=SimpleNamespace(get_message_by_id=AsyncMock())
        with patch.object(e,'Chats',chats),patch.object(e,'ChatMessages',stored):
            result=await self.a.action({'chat_id':'c','messages':[{'id':'m','content':''}]},{'id':'u'},AsyncMock(),AsyncMock())
        self.assertIn('error',result)
        stored.get_message_by_id.assert_not_awaited()

    async def test_clicked_message_and_input_unchanged(self):
        body={'id':'first','messages':[{'id':'first','content':TABLE},{'id':'second','content':'표 없음'}]}
        call=AsyncMock(return_value={'ok':True})
        with contextlib.redirect_stdout(io.StringIO()):
            result=await self.a.action(body,{},None,call)
        self.assertIsNone(result)
        self.assertEqual(body['messages'][1]['content'],'표 없음')
        call.assert_awaited_once()

    async def test_download_error_and_timeout_no_success(self):
        for response in ({'error':'disconnected'},None):
            events=[]
            async def emit(event):events.append(event)
            with contextlib.redirect_stdout(io.StringIO()):
                result=await self.a.action({'messages':[{'content':TABLE}]},{},emit,AsyncMock(return_value=response))
            self.assertIn('error',result)
            self.assertFalse(any(x['data'].get('type')=='success' for x in events))
        with contextlib.redirect_stdout(io.StringIO()):
            result=await self.a.action({'messages':[{'content':TABLE}]},{},AsyncMock(),AsyncMock(side_effect=asyncio.TimeoutError))
        self.assertIn('초과',result['error'])

    async def test_missing_browser_and_tables(self):
        self.assertIn('error',await self.a.action({'messages':[]},{},None,None))
        call=AsyncMock()
        self.assertIn('error',await self.a.action({'messages':[{'content':'본문만'}]},{},None,call))
        call.assert_not_awaited()

    async def test_concurrent_exports_separate_bytes_and_worker(self):
        calls=[AsyncMock(return_value={'ok':True}),AsyncMock(return_value={'ok':True})]
        thread_ids=[]
        original=e.Action._build_workbook
        def build(obj,*args):
            thread_ids.append(threading.get_ident())
            return original(obj,*args)
        with patch.object(e.Action,'_build_workbook',build),contextlib.redirect_stdout(io.StringIO()):
            results=await asyncio.gather(*(self.a.action({'messages':[{'content':TABLE.replace('가나다',value)}]}, {},None,call) for value,call in zip(['FIRST','SECOND'],calls)))
        self.assertEqual(results,[None,None])
        self.assertTrue(all(t!=threading.get_ident() for t in thread_ids))
        self.assertNotEqual(calls[0].call_args.args[0]['data']['code'],calls[1].call_args.args[0]['data']['code'])

    async def test_all_messages_scope(self):
        self.a.valves.EXPORT_SCOPE='all_messages'
        captured=[]
        original=e.Action._build_workbook
        def build(obj,tables,names):
            captured.append(len(tables))
            return original(obj,tables,names)
        with patch.object(e.Action,'_build_workbook',build),contextlib.redirect_stdout(io.StringIO()):
            result=await self.a.action({'messages':[{'content':TABLE},{'content':TABLE}]},{},None,AsyncMock(return_value={'ok':True}))
        self.assertIsNone(result)
        self.assertEqual(captured,[2])

    async def test_ai_timeout_falls_back_to_download(self):
        self.a.valves.TITLE_SOURCE='ai_generated'
        users=SimpleNamespace(get_user_by_id=AsyncMock(return_value=SimpleNamespace(id='u')))
        generate=AsyncMock(side_effect=asyncio.TimeoutError)
        call=AsyncMock(return_value={'ok':True})
        with patch.object(e,'Users',users),patch.object(e,'generate_chat_completion',generate),contextlib.redirect_stdout(io.StringIO()):
            result=await self.a.action({'messages':[{'content':TABLE}]},{'id':'u'},AsyncMock(),call,object())
        self.assertIsNone(result)
        call.assert_awaited_once()

    async def test_javascript_ack_and_xlsx_bytes(self):
        node=shutil.which('node')
        if not node:self.skipTest('Node required for JS contract')
        call=AsyncMock(return_value={'ok':True})
        with contextlib.redirect_stdout(io.StringIO()):
            await self.a.action({'messages':[{'content':TABLE}]},{},None,call)
        code=call.call_args.args[0]['data']['code']
        harness="""
        let clicks=0;
        global.document={createElement:()=>({click:()=>clicks++,remove:()=>{}}),body:{appendChild:()=>{}}};
        global.setTimeout=fn=>fn();
        const AsyncFunction=Object.getPrototypeOf(async function(){}).constructor;
        new AsyncFunction(JSON.parse(process.argv[1]))().then(result=>{
            if(!result || result.ok!==true || clicks!==1)process.exit(1);
        }).catch(()=>process.exit(2));
        """
        subprocess.run([node,'-e',harness,json.dumps(code)],check=True,capture_output=True)

if __name__=='__main__':unittest.main()
