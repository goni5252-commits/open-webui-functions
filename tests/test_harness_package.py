import copy
import importlib.util
import json
from pathlib import Path
import subprocess
import tempfile
import unittest
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1] / 'harness/open-terminal'

def load(name):
    spec = importlib.util.spec_from_file_location('harness_'+name, ROOT/'scripts'/f'{name}.py')
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module

fetch = load('fetch_design')
builder = load('build_pptx')
validator = load('validate_pptx')


class HarnessPackageTests(unittest.TestCase):
    def test_fetch_isolation_success_and_provenance(self):
        with tempfile.TemporaryDirectory() as directory:
            original = Path(directory)/'DESIGN.md'
            original.write_text('keep')
            def runner(argv, **kw):
                self.assertEqual(argv, ['npx','-y','getdesign@latest','add','vercel'])
                self.assertEqual(kw['timeout'], 120)
                (kw['cwd']/'DESIGN.md').write_text('# Design\nColors: #ffffff')
                return subprocess.CompletedProcess(argv, 0)
            a, b = [fetch.fetch('vercel', directory, runner) for _ in range(2)]
            self.assertNotEqual(a['path'], b['path'])
            self.assertEqual(original.read_text(), 'keep')
            self.assertEqual(len(a['sha256']), 64)
            self.assertTrue(Path(a['path']).with_name('source.json').exists())

    def test_fetch_invalid_input_never_runs(self):
        for slug in ('--help','../a','x;id','https://x','a\nb','a'*81):
            with self.assertRaises(ValueError):
                fetch.fetch(slug, '/unused', lambda *a, **k: self.fail('executed'))

    def test_fetch_empty_html_failure_timeout(self):
        for mode in ('empty', 'whitespace', 'html', 'failed', 'timeout'):
            with self.subTest(mode=mode), tempfile.TemporaryDirectory() as directory:
                def runner(argv, **kw):
                    if mode == 'timeout': raise subprocess.TimeoutExpired(argv, 120)
                    if mode in ('html', 'whitespace'):
                        (kw['cwd']/'DESIGN.md').write_text('<html>error' if mode == 'html' else '  ')
                    return subprocess.CompletedProcess(argv, 1 if mode == 'failed' else 0)
                with self.assertRaises((ValueError, RuntimeError, subprocess.TimeoutExpired)):
                    fetch.fetch('vercel', directory, runner)

    def test_builder_creates_editable_korean_and_validator_flags_bounds(self):
        from pptx import Presentation
        theme = json.loads((ROOT/'examples/theme.json').read_text())
        slides = json.loads((ROOT/'examples/slides.json').read_text())
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)/'demo.pptx'
            builder.build(theme, slides, path)
            report = validator.validate(path)
            self.assertTrue(report['ok'], report)
            self.assertEqual(report['slides'], 3)
            self.assertEqual(report['visual_check'], 'not_performed')
            prs = Presentation(path)
            text = '\n'.join(shape.text for slide in prs.slides for shape in slide.shapes if shape.has_text_frame)
            self.assertIn('인공지능 활용 연수', text)
            with self.assertRaises(FileExistsError): builder.build(theme, slides, path)
            prs.slides[0].shapes[0].left = -100
            prs.save(path)
            self.assertFalse(validator.validate(path)['ok'])

    def test_builder_rejects_bad_theme_and_crowded_content(self):
        theme = json.loads((ROOT/'examples/theme.json').read_text())
        slides = json.loads((ROOT/'examples/slides.json').read_text())
        wrong = copy.deepcopy(theme)
        wrong['colors']['accent'] = 'not-a-color'
        with self.assertRaises(ValueError): builder.validate_inputs(wrong, slides)
        slides[1]['body'] = ['too many']*7
        with self.assertRaises(ValueError): builder.validate_inputs(theme, slides)
