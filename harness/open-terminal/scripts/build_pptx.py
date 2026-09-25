"""Small editable PPTX builder; advanced layouts may use separate task scripts."""
import argparse
import json
from pathlib import Path
import re


def validate_inputs(theme, slides):
    if not isinstance(theme, dict) or not isinstance(slides, list) or not 1 <= len(slides) <= 100:
        raise ValueError('Expected theme object and 1–100 slides')
    for key in ('background', 'surface', 'text', 'accent'):
        if not re.fullmatch(r'[0-9a-fA-F]{6}', theme['colors'][key]):
            raise ValueError(f'Invalid RGB color: {key}')
    for key in ('title', 'body'):
        if not isinstance(theme['fonts'][key], str) or not theme['fonts'][key].strip():
            raise ValueError('Specify actual installed font families')
        value = theme['font_sizes_pt'][key]
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not 12 <= value <= 60:
            raise ValueError('Font sizes must be 12–60pt')
    size = theme['slide_size']
    if not all(isinstance(size[k], (int, float)) and 5 <= size[k] <= 30 for k in ('width', 'height')):
        raise ValueError('Slide dimensions must be 5–30 inches')
    margin = theme['margins_inches']
    if not isinstance(margin, (int, float)) or not 0.25 <= margin <= 1.5 or 2*margin+2 >= size['height']:
        raise ValueError('Invalid slide margin')
    for slide in slides:
        if slide.get('layout') not in ('title', 'bullets', 'cards'):
            raise ValueError('Supported layouts: title, bullets, cards')
        if not isinstance(slide.get('title'), str) or not 1 <= len(slide['title']) <= 120:
            raise ValueError('Use a nonempty short slide title')
        if slide['layout'] == 'cards':
            cards = slide.get('cards', [])
            if not isinstance(cards, list) or not 1 <= len(cards) <= 4:
                raise ValueError('Use 1–4 cards per slide')
            if any(not isinstance(c, dict) or any(not isinstance(c.get(k), str) or len(c[k]) > 300 for k in ('title', 'body')) for c in cards):
                raise ValueError('Each card needs short title/body strings')
        else:
            body = slide.get('body', [])
            if not isinstance(body, list) or len(body) > 6 or any(not isinstance(t, str) or len(t) > 300 for t in body):
                raise ValueError('Use up to 6 short body strings; split crowded slides')


def build(theme, slides, output):
    validate_inputs(theme, slides)
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    prs = Presentation()
    width, height = theme['slide_size']['width'], theme['slide_size']['height']
    prs.slide_width, prs.slide_height = Inches(width), Inches(height)
    margin = theme['margins_inches']
    colors = theme['colors']

    def text(slide, content, x, y, w, h, role='body', color='text'):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(0.03)
        for i, line in enumerate(content):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = line
            p.font.name = theme['fonts'][role]
            p.font.size = Pt(theme['font_sizes_pt'][role])
            p.font.bold = role == 'title'
            p.font.color.rgb = RGBColor.from_string(colors[color])
            p.space_after = Pt(12)
        return box

    for number, item in enumerate(slides, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(colors['background'])
        cover = item['layout'] == 'title'
        title_y = height*0.27 if cover else margin
        text(slide, [item['title']], margin, title_y, width-2*margin, 1.2, 'title')
        y = title_y+1.45
        available = height-y-margin-0.25
        if item['layout'] == 'cards':
            cards = item['cards']
            if available < 1.4*((len(cards)+1)//2):
                raise ValueError('Card layout too short: increase slide height or reduce margin/card count')
            gap = 0.25
            cols = min(2, len(cards))
            rows = (len(cards)+cols-1)//cols
            cw = (width-2*margin-gap*(cols-1))/cols
            ch = (available-gap*(rows-1))/rows
            for i, card in enumerate(cards):
                x, cy = margin+(i%cols)*(cw+gap), y+(i//cols)*(ch+gap)
                kind = MSO_SHAPE.ROUNDED_RECTANGLE if theme.get('card_style', {}).get('rounded') else MSO_SHAPE.RECTANGLE
                shape = slide.shapes.add_shape(kind, Inches(x), Inches(cy), Inches(cw), Inches(ch))
                shape.fill.solid()
                shape.fill.fore_color.rgb = RGBColor.from_string(colors['surface'])
                shape.line.fill.background()
                text(slide, [card['title']], x+0.16, cy+0.12, cw-0.32, 0.55, color='accent')
                text(slide, [card['body']], x+0.16, cy+0.75, cw-0.32, max(0.2, ch-0.85))
        else:
            text(slide, item.get('body', []), margin, y, width-2*margin, available)
    path = Path(output)
    path.parent.mkdir(parents=True, exist_ok=True)
    # Never silently replace a prior result.
    with path.open('xb') as stream:
        prs.save(stream)
    return {'path': str(path.resolve()), 'slides': len(slides), 'visual_check': 'not_performed'}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--theme', required=True)
    parser.add_argument('--slides', required=True)
    parser.add_argument('--output', required=True)
    args = parser.parse_args()
    try:
        result = build(json.loads(Path(args.theme).read_text(encoding='utf-8')),
                       json.loads(Path(args.slides).read_text(encoding='utf-8')), args.output)
        print(json.dumps(result, ensure_ascii=False))
    except (OSError, ValueError, KeyError, TypeError, ImportError) as exc:
        parser.exit(1, f'{type(exc).__name__}: {exc}\n')
