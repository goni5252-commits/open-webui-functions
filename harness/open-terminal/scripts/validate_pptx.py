"""Structural and shape-boundary checks, explicitly not a visual rendering check."""
import argparse
import json
from pathlib import Path
from zipfile import ZipFile


def validate(path):
    from pptx import Presentation
    path = Path(path)
    with ZipFile(path) as archive:
        if 'ppt/presentation.xml' not in archive.namelist() or archive.testzip():
            raise ValueError('Invalid PPTX package')
    prs = Presentation(path)
    errors = []
    if not prs.slides:
        errors.append('No slides')
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left+shape.width > prs.slide_width+2 or shape.top+shape.height > prs.slide_height+2:
                errors.append(f'Slide {i}: shape outside slide bounds ({shape.name})')
    return {'ok': not errors, 'slides': len(prs.slides), 'bytes': path.stat().st_size,
            'errors': errors, 'visual_check': 'not_performed',
            'limitations': 'Does not detect text overflow, font substitution or intended/accidental overlap. Render and inspect separately.'}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('file')
    args = parser.parse_args()
    try:
        result = validate(args.file)
        print(json.dumps(result, ensure_ascii=False))
        raise SystemExit(0 if result['ok'] else 1)
    except (OSError, ValueError, ImportError) as exc:
        parser.exit(1, f'{type(exc).__name__}: {exc}\n')
